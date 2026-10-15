from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(r'D:\cxf\code\SageRoute')
OUT = ROOT / 'docs' / 'SageRoute_APP数据接口与数据库说明附录.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(2.1), Cm(2.0)
s.left_margin = s.right_margin = Cm(2.0)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st = d.styles[name]
    st.font.name = 'Microsoft YaHei'
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
d.styles['Normal'].font.size = Pt(10.5)
d.styles['Normal'].paragraph_format.line_spacing = 1.22
d.styles['Normal'].paragraph_format.space_after = Pt(6)
for name,size in [('Title',21),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
    d.styles[name].font.size = Pt(size)
    d.styles[name].font.bold = True
    d.styles[name].paragraph_format.space_before = Pt(12)
    d.styles[name].paragraph_format.space_after = Pt(7)
d.core_properties.title = 'SageRoute APP 数据接口与数据库说明附录'
d.core_properties.subject = '数据接口 路线规划机制 数据库字段'
d.core_properties.author = 'SageRoute'
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def p(t): return d.add_paragraph(t)
def h(t,l=2): return d.add_heading(t,l)
def table(head,rows,widths=None):
    t=d.add_table(rows=1, cols=len(head))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    if widths is None: widths=[17/len(head)]*len(head)
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for i,v in enumerate(head): t.rows[0].cells[i].text=v
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): cells[i].text=str(v)
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');rp.append(no)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');rp.append(repeat)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side,val in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margin.append(e)
            cp.append(margin)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7ECF0');cp.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.line_spacing=1.12
                for run in para.runs:
                    run.font.size=Pt(9)
                    run.bold=(ri==0)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def fields(title,rows,sql=False):
    h(title,3)
    table(['字段名','类型与映射','含义及处理规则'],rows,[4.3,4.3,8.4])

d.add_paragraph('SageRoute APP 数据说明附录',style='Title')
p('数据接口  路线规划机制  数据库字段')
p('版本日期  2026年10月14日')
p('本附录说明 SageRoute 的数据访问接口、高德路径规划器接入方式及数据库字段。道路级路径由高德服务计算，APP 负责组织用户选择的地点、保持访问顺序、合并路段结果并保存行程。')
p('字段依据分为两类：基础内容表按客户端模型的读写约定说明；article_images、user_routes 和 user_route_waypoints 按仓库内 SQL 迁移说明数据库类型与约束。数据库中是否已部署全部迁移，需与运行环境对应核对。')
h('A 数据接口说明',1)
h('A.1 数据来源与访问结构')
p('Supabase 提供认证、文化内容和用户路线数据；高德提供地图、定位及道路路径规划。昵称、简介、语言和旅行偏好使用 SharedPreferences 保存在本机，成就解锁状态也以本地存储为主。Web 路线草稿使用当前浏览器标签页的 sessionStorage，并按用户及草稿编号区分。')
p('页面通过业务仓储取得强类型模型。各仓储复用 SupabaseTableRepository，通过 Supabase SDK 发起查询或写入；DatabaseService 负责初始化、超时与重试。下列方法是 APP 内部调用接口，不代表项目另建了同名 HTTP 服务。Supabase 表访问由 SDK 转换为 Data API 请求。')
h('A.2 用户认证接口')
table(['方法','输入','返回与行为'],[
('signUp','email、password、nickname','AuthResponse；调用 Supabase 注册，昵称写入用户 metadata；是否取得 session 取决于邮箱验证配置。'),
('signIn','email、password','AuthResponse；密码登录，存在 session 时按项目会话策略持久化。'),
('signOut','无','退出登录并清理持久化会话。'),
('requestPasswordReset','email','发送密码重置邮件，并指定应用重置页面地址。'),
('resendVerification','email','重新发送注册验证邮件。'),
('updatePassword','password','在有效会话或密码恢复流程中更新密码。')],[4.0,4.1,8.9])
p('认证由 AuthService 封装 Supabase Auth。业务代码通过当前用户 ID 确定用户路线归属；客户端不建立明文密码业务表。认证错误由页面提示，用户 metadata 中的昵称用于展示，不用于判定路线所有权。')
h('A.3 文化内容和地点接口')
table(['仓储','主要方法及参数','结果'],[
('CelebrityRepository','fetchCelebrities；fetchById(id)；searchByName(query)','人物列表或单个人物；单条不存在时为 null。'),
('DynastyRepository','fetchDynasties(limit)','朝代列表。'),
('TopicRepository','fetchTopics(limit)；fetchTopicsByCelebrity(celebrityName)；fetchTopicById(id)；searchByName(query)','主题列表或单条主题；按人物名称筛选。'),
('LocationRepository','fetchLocations(limit)；fetchById(id)；fetchLocationsByTopic(topicName)；fetchArEnabledLocations','地点列表或单条地点；主题筛选使用数据库字段 Topic。'),
('PoiCelebrityRelationRepository','fetchRelations(limit)；fetchRelationsByCelebrity(name)；fetchRelationsByLocation(name)','人物与地点的关系记录。'),
('ArticleRepository','fetchArticles(limit)；fetchById(id)；searchByTitle(query)','文章列表或单篇文章。'),
('ArticleImageRepository','fetchArticleImages(articleId)','图片记录；articleId 可选，可按文章筛选。'),
('IconRepository','fetchIcons；fetchIconMap','类别图标记录或名称到 SVG 路径数据的映射。')],[4.4,7.5,5.1])
p('通用访问层支持指定查询列、limit、等值过滤和 ilike 模糊查询，并将原始行映射为模型。按 ID 查询未命中时，业务仓储通常返回 null；列表查询未命中时返回空列表。人物—地点关系目前按名称查询，不应将该名称关联描述为已经存在的数据库外键。')
p('Location 模型读取 id (PK)，并兼容 id；但 fetchById 当前按 id 发起过滤。Topic 字段读取兼容大小写，主题过滤使用 Topic。接入某一实际数据库时，字段名称必须与该环境一致。')
h('A.4 高德路径规划接口')
p('统一接口 AmapGateway.fetchRoute(request) 接收 AmapRouteDraftRequest，返回 AmapRouteResult。Android 由 NativeAmapGateway 通过 MethodChannel 调用原生高德 RouteSearch；Web 由 WebAmapGateway 调用地图桥接中的 AMap.Driving 或 AMap.Walking。项目还提供 HTTP 网关类，但当前 createRouteGateway 的入口选择原生或 Web 网关。')
table(['请求字段','类型','说明'],[
('origin、destination','RoutePlace','起终点，包含 id、name、latitude、longitude 等信息。'),
('waypoints','List<RoutePlace>','有序途经点；步行分段请求不携带中间途经点。'),
('transportType','TransportType','driving 或 walking。'),
('preferenceKey','String','路线策略键；当前页面预览传入 shortest。'),
('optimizeOrder','bool','请求模型中保留；当前多地点规划传入 false，不请求重排用户选点顺序。')],[4.5,4.1,8.4])
table(['返回字段','单位或形式','说明'],[
('distanceMeters','整数 米','路线距离。'),('travelDuration','Duration','交通耗时；高德原始时间按秒转换。'),
('polyline','List<List<double>>','统一模型采用 [纬度, 经度] 坐标点，用于绘制路线。'),
('waypointOrder','可空列表','途经点排序信息；当前原生网关返回 null，主流程不依赖其重排。')],[4.5,4.1,8.4])
p('Web 高德请求使用 [经度, 纬度]，Location.coordinates 也采用该顺序；进入统一路线折线模型时使用 [纬度, 经度]。这两种顺序不可直接混用。原生方法名分别为 calculateDriveRoute 和 calculateWalkRoute，通道名称为 com.sageroute/route_planning。')
h('A.5 用户路线接口')
table(['方法','输入','结果及行为'],[
('fetchUserRoutes','当前登录用户','查询 user_routes，按创建时间新到旧排列，并按 sort_order 装配途经点。'),
('fetchRouteById','routeId','返回含途经点的 SavedRoute；不存在时返回 null，访问范围由数据库 RLS 限制。'),
('createRoute','SavedRoute draft','写入父路线，再批量写入途经点，返回生成的路线 UUID。user_id 取当前用户。'),
('updateWaypoints','routeId、waypoints','读取旧途经点后全量替换；插入失败时尝试恢复旧数据。'),
('deleteRoute','routeId','按路线 ID 和当前用户 ID 删除；未删除任何记录时提示不存在或无权删除。')],[4.3,4.3,8.4])
p('createRoute 的两次 REST 写入不构成一个数据库事务。途经点插入失败时，仓储尝试删除已写入的父路线；清理失败会记录日志。编辑也使用失败恢复机制，而非原子事务。路线列表和详情不应据此宣称具有事务级一致性。')
h('A.6 异常处理与数据示例')
p('通用数据库操作单次超时为 25 秒，默认重试 2 次，共最多 3 次尝试；仅对可识别的网络或超时异常重试，退避间隔依次为 400 毫秒、800 毫秒。当前原生和 Web 规划网关的超时为 15 秒。缺少迁移、非空约束冲突、权限不足和网络失败分别转为可供页面展示的错误。')
p('新建路线写入结构如下。标识符及数值为说明用示例，不对应真实账号或行程。父路线 id 由数据库生成，途经点通过 route_id 关联，sort_order 从 0 开始。')
table(['对象','关键字段示例'],[
('user_routes','user_id=<当前用户UUID>；title=文化主题半日路线；date_range=2026-10-14；duration=约3小时；distance=约8公里；figure_id=1；figure_name=苏轼'),
('第一个途经点','route_id=<新路线UUID>；sort_order=0；name=地点甲；latitude=30.25；longitude=120.15；visit_duration_min=30；transport_to_next=walking'),
('最后一个途经点','route_id=<新路线UUID>；sort_order=1；name=地点乙；latitude=30.26；longitude=120.16；visit_duration_min=45；不发送 transport_to_next')],[3.5,13.5])
h('B 路线规划机制说明',1)
h('B.1 高德路径规划器的使用方式')
p('APP 当前使用高德路径规划器计算道路级路线，没有自行构建道路图或实现道路搜索算法。Android 驾车调用 RouteSearch.DriveRouteQuery 和 calculateDriveRouteAsyn，步行调用对应的步行查询；Web 使用 AMap.Driving.search 与 AMap.Walking.search。返回方案中取首条路径，提取距离、耗时和各步骤坐标。')
p('Android 将 shortest 映射到 DRIVING_SINGLE_SHORTEST，其余键映射到 DRIVING_SINGLE_DEFAULT；Web 将 shortest 映射到 LEAST_DISTANCE，其余键映射到 LEAST_TIME。以上是道路路径策略，不代表对多个文化地点的访问顺序进行全局优化。当前页面预览统一使用 shortest。')
p('高德官方 Android 驾车规划指南和 JavaScript 路线规划指南公开了起终点、途经点、策略参数及路线结果的调用方式。所查文档未披露底层搜索或优化算法，因此本文不将高德内部方法指定为 A*、遗传算法或模拟退火。公开文档的策略名称仅说明路线选择偏好。')
h('B.2 用户选点与分段规划')
p('创建路线页面通过 RoutePreviewCoordinator 获取用户已选地点及各相邻路段交通方式。系统保留用户排列顺序，不自动删去地点，也不重排地点。至少两个地点才能发起有效预览；各段交通方式数量须等于地点数量减一。')
p('页面预览按相邻地点逐段规划，可在一条路线内混合步行和驾车。底层 MultiStopRoutePlanner 同时支持完整的同交通方式地点序列：步行逐对拆分，驾车每次至多提交 6 个途经点，即包括起终点共 8 个地点，下一段复用上一段终点。该上限是项目采用的拆分规则，与官方 Android 指南一致；Web 官方指南允许的途经点数量可能不同。')
p('地点经纬度须为有限值，纬度在 −90 至 90 范围，经度在 −180 至 180 范围，并排除 (0,0)。路段返回折线必须包含至少两个可绘制坐标点；失败时标明出错路段，避免以不完整路线作为正常预览结果。')
h('B.3 路线合并与时间统计')
p('系统累加各段距离和交通耗时，并依次拼接折线。无效坐标点被忽略，相邻段的重复边界点被去除，以保持路线连续。统一距离采用米；交通耗时从秒转换为 Duration；地点停留时间采用分钟。')
p('RoutePlanningEngine 另提供总时长与时间约束检查：总时长 = 交通耗时 + 各地点预计游览时长，游览时长缺失时按 0 计。总时长不大于 timeWindow 时通过检查。该引擎存在于代码中，但当前创建页面预览直接使用 RoutePreviewCoordinator，不能把引擎的所有约束检查都视为页面已执行。')
h('B.4 预览更新和算法简述')
p('RoutePreviewCoordinator 默认采用 240 毫秒防抖；用户修改地点或交通方式后发起新的规划。通过递增的请求代次识别过期结果，旧结果不覆盖最新预览。状态分为地点不足、规划中、成功和失败，保存入口据状态决定是否可用或是否需要确认失败情况。')
p('启发式搜索、遗传算法和模拟退火均可用于一般路线优化，但当前项目未实现遗传算法或模拟退火，也未发现 A* 等启发式搜索主流程。place_set_generator.dart 保留了按距离、走廊范围、椭圆绕行范围及地点类别选点的启发式策略，当前页面未调用这些策略。本文仅作简要说明，实际路径计算仍归于高德路径规划器。')
h('B.5 高德公开资料',2)
p('以下资料于2026年10月14日查阅，用于说明公开接口能力及策略参数，不用于推断底层算法。')
p('高德开放平台  Android 地图 SDK 驾车出行路线规划\nhttps://lbs.amap.com/api/android-sdk/guide/route-plan/drive')
p('高德开放平台  JavaScript API 路线规划\nhttps://lbs.amap.com/api/javascript-api/guide/services/navigation')
h('C 数据库结构与字段说明',1)
h('C.1 数据组织与关联关系')
p('基础数据包括人物、朝代、主题、地点、人物地点关系、文章、文章图片和类别图标。人物、主题、地点之间的部分关联采用名称或文本匹配；数据库字段不能仅凭客户端模型就认定存在外键。已由迁移明确的关系为：article_images.article_id 指向 Article.id，user_routes.user_id 指向 auth.users，user_route_waypoints.route_id 指向 user_routes.id，均支持父记录删除时级联删除。')
p('下文基础表中的“类型与映射”表示 Dart 读取类型及其空值处理，不等同于 PostgreSQL 实际列类型或 NOT NULL 约束。路线及文章图片表采用 SQL 类型，明确说明是否可空和数据库默认值。字符串数组在数据库中究竟为数组列还是 JSON，需要以实际表结构为准。')
h('C.2 文化内容基础表')
fields('C.2.1 Celebrity 历史人物',[
('id','int → id','人物编号；无法读取数值时模型使用 0。'),('name','String → name','人物姓名；缺失时为空字符串。'),('dynasty','String → dynasty','朝代名称；缺失时为空字符串。'),('bio_short','String → bioShort','人物简述。'),('bio_full 或 bio_ful','String → bioFull','读取优先 bio_ful，兼容 bio_full；模型写出 bio_full。'),('avatar_url','String → avatarUrl','头像资源地址。'),('topic','List<String> → topic','关联主题名称列表；非列表值按空列表处理。')])
fields('C.2.2 Dynasty 朝代', [('id','int → id','朝代编号；缺失时为 0。'),('dynasty','String → dynasty','朝代名称；缺失时为空字符串。')])
fields('C.2.3 Topic 文化主题',[
('id','int → id','主题编号。'),('created_at','DateTime? → createdAt','创建时间；解析失败时为 null。'),('celebrity','String? → celebrity','关联人物名称，按该名称查询人物主题。'),('name','String → name','主题名称。'),('description','String? → description','主题说明，可在模型中为空。')])
fields('C.2.4 Location 地点',[
('id (PK) 或 id','int → id','模型优先读取 id (PK)，兼容 id；写出 id (PK)。按 ID 查询方法当前过滤 id。'),
('name_modern','String → nameModern','现代地名。'),('name_ancient','String? → nameAncient','古代地名。'),('description','String? → description','地点介绍。'),
('average_visit_duration_min','int? → averageVisitDurationMin','预计游览时长，单位分钟。'),('address','String? → address','地点地址。'),('open_time','String? → openTime','开放时间的文本表示，模型未强制时间格式。'),('close_time','String? → closeTime','关闭时间的文本表示。'),('is_ar_enabled','bool → isArEnabled','只有原值为 true 时开启；其他值按 false 处理。'),('Topic 或 topic','String? → topic','主题名称；读取优先 Topic，查询及写出使用 Topic。'),('categories','List<String> → categories','类别列表；也兼容以逗号、顿号或分号分隔的文本。'),('coordinates','List<double> → coordinates','[经度, 纬度]；不足两个值不作为有效路线地点。')])
fields('C.2.5 poi_celebrity_relatian 人物地点关系',[
('id','int → id','关系记录编号；表名沿用原有 relatian 拼写。'),('created_at','DateTime? → createdAt','创建时间。'),('location_name','String? → locationName','关联地点名称。'),('celebrity_name','String? → celebrityName','关联人物名称。'),('relation_type','String? → relationType','关系类型文本。'),('weight','int? → weight','关联权重；模型未定义统一范围或评分公式。')])
fields('C.2.6 Article 文章',[
('id','int → id','文章编号。'),('created_at','DateTime? → createdAt','创建时间。'),('topic','String → topic','文章主题文本。'),('title','String → title','文章标题。'),('summary','String → summary','文章摘要。'),('content','String → content','文章正文。'),('cover_image_url','String → coverImageUrl','封面地址；读取兼容 coverImageUrl。')])
fields('C.2.7 Icons 类别图标',[
('name','String → name','地点类别名称。'),('svg 或 SVG','String → svg','SVG path 的 d 路径数据；读取兼容 SVG，写出 svg，并非必然是完整 SVG 文档。')])
h('C.3 文章图片表')
fields('C.3.1 article_images',[
('id','bigint → id','主键，非空，generated by default as identity。'),('article_id','bigint → articleId','非空；外键关联 Article.id，删除文章时级联删除。'),('role','text → role','非空；cover、featured、poster 三选一。模型未知值回退为 cover。'),('image_url','text → url','非空且去空白后长度大于 0；保存图片地址，不保存二进制图片。'),('alt_text','text → altText','可空；替代文本。'),('source','text → source','可空；图片来源。'),('photographer_name','text → photographerName','可空；摄影者姓名。'),('attribution_url','text → attributionUrl','可空；来源或署名链接。'),('sort_order','integer → sortOrder','非空，默认 0，且不小于 0。'),('created_at','timestamptz → createdAt','非空，默认 UTC 当前时间。')],True)
p('组合唯一约束为 (article_id, role)，每篇文章每种图片角色最多一条记录；索引为 (article_id, sort_order)。启用 RLS，并明确允许 anon 和 authenticated 查询；迁移未为客户端授予图片表写入权限。')
h('C.4 用户路线数据表')
fields('C.4.1 user_routes',[
('id','uuid → id','主键，非空，默认 gen_random_uuid()；Dart 使用 String。'),('user_id','uuid','非空；关联 auth.users，删除用户时级联删除。客户端取当前用户 ID。'),('title','text → title','非空，去首尾空白后长度大于 0。'),('date_range','text → dateRange','可空；日期或日期范围的展示文本。'),('duration','text → duration','可空；时长展示文本，不是数值时长字段。'),('distance','text → distance','可空；距离展示文本，不是米制数值字段。'),('figure_id','bigint → figureId','可空；人物编号，当前迁移未声明指向 Celebrity 的外键。'),('figure_name','text → figureName','可空；人物名称快照。'),('route_snapshot','jsonb','非空，默认 {}；预留路线快照，当前 SavedRoute.toMap 不写入此字段。'),('created_at','timestamptz','非空，默认 UTC 当前时间，用于列表排序。'),('updated_at','timestamptz','非空，默认 UTC 当前时间；更新父路线时由触发器刷新。')],True)
p('索引为 (user_id, created_at desc)。SavedRoute 模型主要读取展示字段及途经点，不包含 user_id、route_snapshot 和时间戳属性；这些字段仍属于表结构。仅更新途经点不会自动触发父路线的 updated_at 更新。')
fields('C.4.2 user_route_waypoints',[
('id','bigint','主键，非空，generated always as identity。'),('route_id','uuid','非空；关联 user_routes.id，父路线删除时级联删除。'),('sort_order','integer','非空且不小于 0；客户端按列表索引写入，读取时升序排列。'),('name','text → name','非空；途经点名称快照。'),('latitude','double precision → latitude','非空；纬度。'),('longitude','double precision → longitude','非空；经度。'),('location_id','bigint → locationId','可空；目录地点编号，未声明 Location 外键；无有效目录 ID 时为空。'),('location_name','text','可空；兼容描述字段，当前 RouteWaypoint.toMap 不写入。'),('visit_duration_min','integer → visitDurationMin','可空；停留时长，分钟，当前客户端使用该字段。'),('average_visit_duration_min','integer','可空且非负；兼容字段，当前模型不直接读写。'),('categories','text → categories','可空；地点分类快照，区别于 Location 模型的分类列表。'),('topic','text','可空；预留主题字段，当前模型不直接读写。'),('transport_to_next','text → transportToNext','可空；driving 或 walking。最后一点无下一段，读取时规范为 null。'),('transport_type','text','非空，默认 driving；driving 或 walking。兼容字段，非当前分段交通方式的主要来源。'),('created_at','timestamptz','非空，默认 UTC 当前时间。')],True)
p('组合唯一约束为 (route_id, sort_order)，索引采用相同字段。客户端对末点不发送 transport_to_next，以兼容旧库默认值；仓库中的后续迁移移除该字段的非空限制。坐标范围检查在规划层完成，表迁移未给 latitude、longitude 声明范围 CHECK。visit_duration_min 也未在此迁移中声明非负 CHECK。')
h('C.5 本地字段与认证数据')
table(['存储项','形式或键','说明'],[
('昵称','profile.nickname','本地覆盖；空字符串清除覆盖并回退到账号昵称。'),('简介','profile.bio','本地个人简介。'),('界面语言','profile.language','默认 zh-Hans；当前仅支持简体中文。'),('旅行偏好','profile.travel_preferences','JSON：dynasties、themes、pace、transport。'),('Web 路线草稿','sageroute.draft.<userId>.<draftId>','sessionStorage 中的 JSON，version=1；损坏或版本不符时清理。'),('成就','本地成就定义与解锁存储','成就定义来自项目数据，解锁状态使用本地存储，不描述为已上线的云端成就表。')],[3.1,7.0,6.9])
p('auth.users 由 Supabase Auth 管理，user_routes.user_id 引用其用户标识。用户资料与旅行偏好当前以本地实现为准；docs/supabase_profile_schema.md 中的 profiles、云端成就及探索记录属于方案内容，不纳入本附录的已实现业务表字段字典。')
h('C.6 权限和数据完整性')
p('用户路线与途经点表启用 RLS，对 anon 撤销表权限，对 authenticated 授予查询、插入、更新和删除权限。路线策略以 auth.uid() = user_id 作为所有权条件；途经点通过所属路线核验用户。USING 与 WITH CHECK 同时限制可操作记录与写入后归属。')
p('路线和途经点的父子关系由外键及级联删除保证，途经点顺序由组合唯一约束保证。保存与全量替换途经点采用客户端补偿清理或恢复，存在多请求之间的时间窗口；批量插入及所有权约束仍由数据库执行。基础内容表的读取权限与具体列约束，应以各运行环境实际 schema 和策略为准。')
p('字段追溯依据：lib/models 下的对应记录模型、lib/data 下的业务仓储、lib/services/auth_service.dart、lib/route_planning 下的规划组件，以及 supabase/migrations 中的文章图片和用户路线迁移。路线主流程对应 step3_map.dart、route_preview_coordinator.dart、multi_stop_route_planner.dart、原生 RoutePlanningHandler.kt 和 web/sage_maps.js。')

foot=s.footer.paragraphs[0]
foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=foot.add_run('SageRoute 数据说明附录  ·  ');run.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
OUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUT)
print(OUT)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables))
