from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from copy import deepcopy

ROOT = Path(r'D:\cxf\code\SageRoute')
OUT = ROOT / 'docs' / 'SageRoute_APP数据说明附录_调研报告格式.docx'
REFERENCE=Path(r'C:\Users\importme\xwechat_files\wxid_9vkc2ceeqrms12_ca53\msg\file\2026-10\东坡杭行调研报告(1).docx')
d = Document(REFERENCE)
reference_section=deepcopy(d.sections[2]._sectPr)
header_part=d.sections[2].header.part
header_rel=d.part.relate_to(header_part,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/header')
header_ref=OxmlElement('w:headerReference'); header_ref.set(qn('w:type'),'default'); header_ref.set(qn('r:id'),header_rel)
reference_section.insert(0,header_ref)
for item in list(d._element.body): d._element.body.remove(item)
d._element.body.append(reference_section)
s = d.sections[0]
s.page_width = Cm(21.59); s.page_height = Cm(27.94)
s.top_margin = s.bottom_margin = Cm(2.54)
s.top_margin=Cm(4.45)
s.left_margin = Cm(2.7993); s.right_margin = Cm(2.6)
s.header_distance = s.footer_distance = Cm(1.5)

def font(run, size=10.5, bold=False):
    run.font.name = '宋体'; run.font.size = Pt(size); run.font.bold = bold
    rpr = run._element.get_or_add_rPr(); rf = rpr.get_or_add_rFonts()
    for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k), '宋体')
    lang = rpr.find(qn('w:lang'))
    if lang is None: lang = OxmlElement('w:lang'); rpr.append(lang)
    lang.set(qn('w:eastAsia'),'zh-CN')

def para(text='', size=10.5, bold=False, align=None, before=0, after=0, line=1.5, first=True):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.line_spacing=line
    if first: p.paragraph_format.first_line_indent=Pt(size*2)
    p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if text:
        r=p.add_run(text); font(r,size,bold)
    return p

def heading(text, level=1):
    if level==1: size,before,after=16,12,8
    elif level==2: size,before,after=14,10,6
    else: size,before,after=11,7,4
    p=para(text,size,bold=True,before=before,after=after,line=1.25,first=False)
    p.paragraph_format.keep_with_next=True
    outline=OxmlElement('w:outlineLvl'); outline.set(qn('w:val'),str(level-1)); p._p.get_or_add_pPr().append(outline)
    return p

def set_cell(cell, text, header=False):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    r=p.add_run(str(text)); font(r,9.2,header)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcPr=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for side,val in [('top','45'),('bottom','45'),('left','70'),('right','70')]:
        e=OxmlElement('w:'+side); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); mar.append(e)
    tcPr.append(mar)
    if header:
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'D9EAF7'); tcPr.append(sh)

def tbl(headers, rows, widths):
    t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=[w*16.19/sum(widths) for w in widths]
    for i,w in enumerate(widths): t.columns[i].width=Cm(w)
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i],h,True)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell(cells[i],v)
    # Repeat header and light borders
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells): cell.width=Cm(widths[ci])
        trPr=row._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
        if ri==0: trPr.append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'A6A6A6'); borders.append(e)
    t._tbl.tblPr.append(borders)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def field_table(title, rows):
    heading(title,3); tbl(['字段','类型或映射','功能说明'],rows,[4.2,4.1,8.7])

# Default styles and page system derived from the supplied research report.
for st in d.styles:
    st.font.name='宋体'; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),'宋体')
d.styles['Normal'].font.size=Pt(10.5)

p=para('SageRoute APP 数据接口与数据库说明附录',16,True,WD_ALIGN_PARAGRAPH.CENTER,after=8,line=1.25,first=False)
para('本附录说明项目中的数据接口、路线规划器接入方式和数据库字段。内容依据 Flutter 代码、原生 Android 路线处理、Web 地图桥接及 Supabase 迁移文件整理。',10.5,False,WD_ALIGN_PARAGRAPH.JUSTIFY,after=6)
para('高德路径规划器负责道路级路线计算；APP 负责组织地点、保持用户选择顺序、拆分请求、合并结果和保存路线。',10.5,False,WD_ALIGN_PARAGRAPH.JUSTIFY,after=10)

heading('一、数据接口说明',1)
heading('（一）数据来源与访问结构',2)
para('Supabase 存储人物、朝代、主题、地点、文章及用户路线。高德提供地图、定位和步行、驾车路径规划。本地 SharedPreferences 保存昵称、简介、语言、旅行偏好和成就状态；Web 路线草稿保存在当前标签页的 sessionStorage。页面通过业务仓储访问数据，SupabaseTableRepository 统一执行查询、写入和过滤。',first=True)
heading('（二）文化内容和地点接口',2)
tbl(['仓储','主要方法','功能'],[
('CelebrityRepository','fetchCelebrities；fetchById；searchByName','读取、查询历史人物。'),
('DynastyRepository','fetchDynasties','读取朝代。'),
('TopicRepository','fetchTopics；fetchTopicsByCelebrity；fetchTopicById；searchByName','读取、筛选文化主题。'),
('LocationRepository','fetchLocations；fetchById；fetchLocationsByTopic；fetchArEnabledLocations','读取地点、主题地点和 AR 地点。'),
('PoiCelebrityRelationRepository','fetchRelationsByCelebrity；fetchRelationsByLocation','读取人物与地点关系。'),
('ArticleRepository','fetchArticles；fetchById；searchByTitle','读取和检索文章。'),
('ArticleImageRepository','fetchArticleImages(articleId)','读取文章图片及署名信息。'),
('IconRepository','fetchIcons；fetchIconMap','读取类别图标和 SVG 路径。')],[4.8,7.5,4.7])
para('查询结果映射为 lib/models 下的强类型对象。ID 未命中时单条查询返回 null，列表查询返回空列表。地点坐标在 Location 模型中按 [经度，纬度] 保存。',first=True)
heading('（三）高德路径规划接口',2)
para('AmapGateway.fetchRoute 接收起点、终点、途经点、交通方式、路线策略和是否优化顺序，返回距离、耗时、路线折线和途经点顺序。Android 使用 NativeAmapGateway，通过 MethodChannel 调用高德 RouteSearch；Web 使用 WebAmapGateway，调用 AMap.Driving 或 AMap.Walking。',first=True)
tbl(['输入或输出','项目字段','说明'],[
('起终点','origin、destination','RoutePlace，包含名称和经纬度。'),
('途经点','waypoints','按用户顺序排列；步行分段不携带中间点。'),
('交通方式','transportType','driving 或 walking。'),
('策略','preferenceKey','当前页面预览使用 shortest。'),
('结果','distanceMeters、travelDuration、polyline','分别为米、Duration 和折线坐标。')],[4.3,5.3,7.4])
para('Web 高德请求使用 [经度，纬度]，统一路线折线使用 [纬度，经度]。原生调用名称为 calculateDriveRoute 和 calculateWalkRoute，通道为 com.sageroute/route_planning。',first=True)
heading('（四）用户路线接口',2)
tbl(['方法','输入','功能'],[
('fetchUserRoutes','当前用户','读取路线并按创建时间倒序装配途经点。'),
('fetchRouteById','routeId','读取单条路线及途经点。'),
('createRoute','SavedRoute','先写 user_routes，再批量写 user_route_waypoints。'),
('updateWaypoints','routeId、waypoints','按 sort_order 全量替换途经点。'),
('deleteRoute','routeId','按路线 ID 和当前用户 ID 删除路线。')],[4.5,4.5,8.0])

heading('二、路线规划机制说明',1)
heading('（一）高德路径规划器',2)
para('项目不自建道路网络，也不在 Flutter 层实现道路搜索。Android 驾车使用 RouteSearch.DriveRouteQuery，步行使用对应查询；Web 使用 AMap.Driving.search 和 AMap.Walking.search。项目从高德返回结果中读取首条路径的距离、耗时和折线。',first=True)
para('Android 将 shortest 映射为 DRIVING_SINGLE_SHORTEST，其他策略使用 DRIVING_SINGLE_DEFAULT；Web 将 shortest 映射为 LEAST_DISTANCE，其他策略使用 LEAST_TIME。该策略控制高德道路方案选择，不负责多个文化地点的访问顺序。',first=True)
para('高德公开文档说明了起终点、途经点、策略参数和结果结构，但未公开底层算法。因此本项目不将高德内部方法表述为 A*、遗传算法或模拟退火。',first=True)
heading('（二）用户选点和分段处理',2)
para('路线创建页保留用户选择的地点顺序，不自动删除或重排地点。地点少于两个时不发起预览；每一段交通方式对应相邻两个地点。步行路线按相邻地点逐段请求，驾车路线按每次最多 6 个途经点拆分，下一段复用上一段终点。',first=True)
para('系统累加各段距离和耗时，拼接折线并删除相邻重复点。路线规划层校验经纬度范围和折线有效性；RoutePlanningEngine 还可将交通耗时与地点游览时长相加，检查 timeWindow。',first=True)
heading('（三）算法简述',2)
para('项目未实现遗传算法和模拟退火，也未发现 A* 主流程。place_set_generator.dart 保留了基于距离、走廊范围、椭圆绕行范围和地点类别的候选点筛选策略，但当前页面以用户顺序和高德路径规划结果为准。',first=True)
heading('（四）预览更新',2)
para('RoutePreviewCoordinator 使用 240 毫秒防抖，在地点或交通方式变化后重新请求，并用请求代次丢弃过期结果。页面显示地点不足、规划中、成功和失败四种状态。',first=True)
heading('（五）高德公开资料',2)
para('高德开放平台 Android 地图 SDK 驾车出行路线规划：\nhttps://lbs.amap.com/api/android-sdk/guide/route-plan/drive\n高德开放平台 JavaScript API 路线规划：\nhttps://lbs.amap.com/api/javascript-api/guide/services/navigation',first=True)

heading('三、数据库结构与字段说明',1)
heading('（一）表关系和说明范围',2)
para('基础表包括 Celebrity、Dynasty、Topic、Location、poi_celebrity_relatian、Article、article_images 和 Icons。已由迁移文件明确的关系为 article_images.article_id → Article.id、user_routes.user_id → auth.users、user_route_waypoints.route_id → user_routes.id；路线父记录删除时，途经点级联删除。基础表字段类型以下客户端模型映射为主，路线和文章图片字段以 SQL 迁移为准。',first=True)
field_table('（二）Celebrity 历史人物',[
('id','int → id','人物编号。'),('name','String → name','人物姓名。'),('dynasty','String → dynasty','朝代名称。'),('bio_short','String → bioShort','人物简述。'),('bio_ful 或 bio_full','String → bioFull','兼容两个字段名，模型写出 bio_full。'),('avatar_url','String → avatarUrl','头像地址。'),('topic','List<String> → topic','关联主题列表。')])
field_table('（三）Dynasty 朝代', [('id','int → id','朝代编号。'),('dynasty','String → dynasty','朝代名称。')])
field_table('（四）Topic 文化主题', [('id','int → id','主题编号。'),('created_at','DateTime? → createdAt','创建时间。'),('celebrity','String? → celebrity','关联人物名称。'),('name','String → name','主题名称。'),('description','String? → description','主题说明。')])
field_table('（五）Location 地点',[
('id (PK) 或 id','int → id','模型兼容两种主键字段；当前 ID 查询过滤 id。'),('name_modern','String → nameModern','现代地名。'),('name_ancient','String? → nameAncient','古代地名。'),('description','String? → description','地点介绍。'),('average_visit_duration_min','int? → averageVisitDurationMin','预计游览时长，分钟。'),('address','String? → address','地点地址。'),('open_time、close_time','String?','开放和关闭时间文本。'),('is_ar_enabled','bool → isArEnabled','AR 开关。'),('Topic 或 topic','String? → topic','主题名称，读取兼容大小写。'),('categories','List<String> → categories','类别列表。'),('coordinates','List<double> → coordinates','[经度，纬度]，至少两个值才可规划。')])
field_table('（六）poi_celebrity_relatian 人物地点关系',[
('id','int','关系编号。'),('created_at','DateTime?','创建时间。'),('location_name','String?','地点名称。'),('celebrity_name','String?','人物名称。'),('relation_type','String?','关系类型。'),('weight','int?','关联权重。')])
field_table('Article 文章',[
('id','int','文章编号。'),('created_at','DateTime?','创建时间。'),('topic','String','文章主题。'),('title','String','文章标题。'),('summary','String','文章摘要。'),('content','String','正文。'),('cover_image_url','String → coverImageUrl','封面地址。')])
field_table('Icons 类别图标',[('name','String','类别名称。'),('svg 或 SVG','String → svg','SVG path 数据。')])
heading('（七）article_images',2)
tbl(['字段','SQL 类型与约束','用途'],[
('id','bigint，identity，主键','图片记录编号。'),('article_id','bigint，非空，外键','关联 Article，级联删除。'),('role','text，cover/featured/poster','图片角色；与 article_id 组合唯一。'),('image_url','text，非空','外部图片地址。'),('alt_text、source、photographer_name、attribution_url','text，可空','替代文本、来源和署名。'),('sort_order','integer，默认 0，非负','文章内排序。'),('created_at','timestamptz，默认 UTC now()','创建时间。')],[4.2,7.0,5.8])
heading('（八）user_routes',2)
tbl(['字段','SQL 类型与约束','用途'],[
('id','uuid，默认 gen_random_uuid()','路线编号。'),('user_id','uuid，非空，引用 auth.users','路线所属用户。'),('title','text，非空','路线标题。'),('date_range、duration、distance','text，可空','日期、时长和距离展示文本。'),('figure_id','bigint，可空','人物编号快照，迁移未声明外键。'),('figure_name','text，可空','人物名称快照。'),('route_snapshot','jsonb，默认 {}','路线快照预留字段。'),('created_at、updated_at','timestamptz，默认 UTC now()','创建和更新时间。')],[4.8,6.0,6.2])
para('user_routes 建有 (user_id, created_at desc) 索引。SavedRoute 当前主要读写标题、日期、时长、距离、人物和途经点，不写 route_snapshot。',first=True)
heading('（九）user_route_waypoints',2)
tbl(['字段','SQL 类型与约束','用途'],[
('id','bigint，identity，主键','途经点编号。'),('route_id','uuid，非空，引用 user_routes','所属路线，级联删除。'),('sort_order','integer，非空且非负','路线顺序；与 route_id 组合唯一。'),('name','text，非空','地点名称快照。'),('latitude、longitude','double precision，非空','纬度和经度。'),('location_id','bigint，可空','目录地点编号，未声明 Location 外键。'),('visit_duration_min','integer，可空','停留时间，分钟。'),('categories、topic','text，可空','地点分类和主题快照。'),('transport_to_next','text，可空，driving/walking','前往下一点的方式，末点为空。'),('transport_type','text，默认 driving','兼容字段，取 driving/walking。'),('created_at','timestamptz，默认 UTC now()','创建时间。')],[4.7,7.0,5.3])
para('兼容字段 location_name 为可空 text，average_visit_duration_min 为可空且非负 integer；当前 RouteWaypoint 模型不直接读写这两个字段。',first=True)
heading('（十）本地数据和权限',2)
tbl(['存储项','键或形式','功能'],[
('昵称、简介','profile.nickname；profile.bio','本地个人资料覆盖。'),('语言','profile.language','默认 zh-Hans。'),('旅行偏好','profile.travel_preferences JSON','保存 dynasties、themes、pace、transport。'),('Web 路线草稿','sageroute.draft.<userId>.<draftId>','sessionStorage，带 version=1。'),('成就','本地成就定义和解锁记录','当前不使用云端成就表。')],[4.2,6.8,6.0])
para('user_routes 和 user_route_waypoints 启用 RLS。authenticated 角色可访问表，策略使用 auth.uid() 校验路线所有权，并通过所属路线校验途经点。anon 角色被撤销路线表权限。',first=True)

# Source header/footer retain background, logo, frame and page numbers.

d.core_properties.title='SageRoute APP 数据接口与数据库说明附录'; d.core_properties.author='SageRoute'
OUT.parent.mkdir(parents=True,exist_ok=True); d.save(OUT)
print(OUT, 'paragraphs',len(d.paragraphs),'tables',len(d.tables))
